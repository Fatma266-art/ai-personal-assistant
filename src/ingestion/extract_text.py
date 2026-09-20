from pathlib import Path
import fitz  # PyMuPDF
from tqdm import tqdm
import logging
from docx import Document
from pptx import Presentation

# =========================
# Setup
# =========================

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent.parent

RAW_DATA_DIR = PROJECT_ROOT / "Data" / "Raw Data"
PROCESSED_DATA_DIR = PROJECT_ROOT / "Data" / "Processed Data"

PROCESSED_DATA_DIR.mkdir(parents=True, exist_ok=True)

# =========================
# PDF Extraction
# =========================

def extract_text_from_pdf(pdf_path: Path) -> str:
    """
    Extract text from PDF while preserving page boundaries.
    Each page starts with a [PAGE X] marker.
    """
    try:
        with fitz.open(pdf_path) as doc:
            pages = []

            for page_number, page in enumerate(doc, start=1):
                text = page.get_text("text")

                if not text.strip():
                    logging.warning(f"Page {page_number} in {pdf_path.name} is empty (might be scanned image)")
                    continue

                page_text = text.strip()
                pages.append(
                    f"[PAGE {page_number}]\n"
                    f"{page_text}"
                )

            return "\n\n".join(pages)

    except Exception as e:
        logging.error(f"Failed to extract from {pdf_path.name}: {e}")
        return ""


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt"}


def extract_text_from_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(parts)
    return f"[PAGE 1]\n{text}" if text.strip() else ""


def extract_text_from_pptx(path: Path) -> str:
    prs = Presentation(str(path))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if texts:
            slides.append(f"[PAGE {i}]\n" + "\n".join(texts))
    return "\n\n".join(slides)


def extract_text(file_path) -> str:
    """Extract text from PDF, DOCX, PPTX or TXT. Returns '' if unsupported."""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    if suffix == ".docx":
        return extract_text_from_docx(path)
    if suffix == ".pptx":
        return extract_text_from_pptx(path)
    if suffix == ".txt":
        text = path.read_text(encoding="utf-8", errors="ignore")
        return f"[PAGE 1]\n{text}" if text.strip() else ""

    logging.warning(f"Unsupported file type: {path.name}")
    return ""


# =========================
# Process PDFs
# =========================

def process_all_pdfs():
    """Process all PDF files in the Raw Data directory."""
    pdf_files = list(RAW_DATA_DIR.glob("*.pdf"))

    if not pdf_files:
        logging.warning("No PDF files found in Raw Data directory.")
        return

    logging.info(f"Found {len(pdf_files)} PDF file(s)")

    for pdf_path in tqdm(pdf_files, desc="Extracting text"):
        logging.info(f"Processing: {pdf_path.name}")

        text = extract_text_from_pdf(pdf_path)

        if text:
            output_path = PROCESSED_DATA_DIR / f"{pdf_path.stem}.txt"

            output_path.write_text(
                text,
                encoding="utf-8"
            )

            logging.info(f"Saved: {output_path.name} | Characters: {len(text)}")

if __name__ == "__main__":
    process_all_pdfs()